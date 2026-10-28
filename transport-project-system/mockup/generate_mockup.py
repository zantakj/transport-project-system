"""
generate_mockup.py — สคริปต์สร้างข้อมูลจำลอง (Mock up) สำหรับระบบบริหารจัดการการขนส่ง

หลักการ: ทุกตัวเลขในไฟล์มาจาก "สมมติฐาน" ในส่วน PARAMETERS ด้านล่างเท่านั้น
ไม่มีการสุ่มตัวเลขแบบไร้ที่มา และใช้ random seed คงที่ → รันซ้ำได้ผลเหมือนเดิมทุกครั้ง (reproducible)

ไฟล์เที่ยววิ่ง (Trips) เก็บเฉพาะ "ข้อมูลดิบที่เกิดขึ้นจริงในการปฏิบัติงาน"
ส่วนต้นทุน ค่าเสื่อม กำไร และคาร์บอน ให้เว็บเป็นผู้คำนวณจากตาราง master
(Vehicles / Routes / Cost_Parameters / Emission_Factors)

วิธีรัน:  python generate_mockup.py   → ได้ไฟล์ Transport_Mockup_v2.xlsx
"""
import math
import random
from datetime import date, datetime, timedelta

import pandas as pd

SEED = 2026
random.seed(SEED)

# =====================================================================
# PARAMETERS (สมมติฐาน) — ทุกค่าต้องอธิบายที่มาในเล่ม / แทนด้วยค่าจริงเมื่อได้ข้อมูลบริษัท
# =====================================================================
START, END = date(2026, 4, 1), date(2026, 9, 30)   # ช่วงข้อมูล 6 เดือน
DEPOT = "BKK"                                        # คลังต้นทาง: กรุงเทพฯ

# ---- ตารางเส้นทาง (Route master) : ระยะทางทางถนนโดยประมาณ กรุงเทพฯ → ตัวเมืองจังหวัด
# ต้องตรวจสอบกับ Google Maps / กรมทางหลวง ก่อนใช้ในเล่ม
ROUTES = [
    # code, ชื่อไทย, ชื่ออังกฤษ, ภาค, ระยะทางเที่ยวเดียว (กม.), ค่าทางด่วนไป-กลับ 4W/6W/10W
    ("SPK", "สมุทรปราการ", "Samut Prakan", "Central", 30, 100, 160, 220),
    ("NPT", "นครปฐม", "Nakhon Pathom", "Central", 56, 0, 0, 0),
    ("AYT", "พระนครศรีอยุธยา", "Phra Nakhon Si Ayutthaya", "Central", 76, 100, 160, 220),
    ("CBI", "ชลบุรี", "Chon Buri", "East", 81, 120, 200, 280),
    ("SRI", "สระบุรี", "Saraburi", "Central", 108, 100, 160, 220),
    ("RYG", "ระยอง", "Rayong", "East", 180, 120, 200, 280),
]

# ---- พารามิเตอร์ตามประเภทรถ
VEHICLE_TYPES = {
    #           จำนวนรถ, คันสำรอง(จอด), km/L, ราคาซื้อ, อายุใช้งาน(ปี), มูลค่าซาก%, เบี้ยประกัน+ภาษี/ปี,
    #           เงินเดือนคนขับ/เดือน, เบี้ยเลี้ยง/เที่ยว, จำนวนยาง, ราคายาง/เส้น, อายุยาง(กม.),
    #           ค่าถ่ายน้ำมันเครื่อง/ครั้ง, ระยะถ่าย(กม.), ค่าซ่อมบำรุงทั่วไป(บาท/กม.),
    #           น้ำหนักบรรทุกสูงสุด(กก.), ค่าบริการเริ่มต้น(บาท/เที่ยว), อัตราค่าขนส่ง(บาท/กม. เที่ยวเดียว)
    "4W":  dict(n=7, standby=1, kmpl=10.0, price=750_000, life=5, salvage=0.10, ins_tax=22_000,
                salary=13_000, allowance=250, tires=4, tire_price=3_500, tire_life=50_000,
                oil_cost=1_800, oil_interval=10_000, repair_km=0.50,
                capacity=1_500, base_fare=1_000, rate_km=24),
    "6W":  dict(n=7, standby=1, kmpl=6.0, price=1_400_000, life=5, salvage=0.10, ins_tax=35_000,
                salary=15_000, allowance=300, tires=6, tire_price=7_000, tire_life=80_000,
                oil_cost=3_500, oil_interval=15_000, repair_km=1.00,
                capacity=5_000, base_fare=1_700, rate_km=35),
    "10W": dict(n=6, standby=0, kmpl=4.0, price=2_600_000, life=5, salvage=0.10, ins_tax=50_000,
                salary=17_000, allowance=350, tires=10, tire_price=9_000, tire_life=100_000,
                oil_cost=6_000, oil_interval=20_000, repair_km=1.50,
                capacity=12_000, base_fare=2_700, rate_km=48),
}

FUEL_TYPE = "Diesel"
# ราคาน้ำมันดีเซลเฉลี่ยรายเดือน (บาท/ลิตร) — สมมติ ต้องแทนด้วยราคาจริงจาก สนพ./ปตท.
DIESEL_PRICE = {4: 32.0, 5: 32.4, 6: 32.9, 7: 32.5, 8: 31.9, 9: 32.2}

# ค่าการปล่อย CO2 ของดีเซล (กก.CO2/ลิตร) — คำนวณตามค่าเริ่มต้นของ IPCC 2006:
#   74,100 kgCO2/TJ × ค่าความร้อน ~36.4 MJ/L (0.0000364 TJ/L) ≈ 2.70 kgCO2/L
# ให้แทนด้วยค่า Emission Factor ฉบับล่าสุดของ อบก. (TGO) และระบุแหล่งอ้างอิงในเล่ม
EF_DIESEL = 2.70

TRIP_PROB_PER_WORKDAY = 0.85   # โอกาสที่รถ active 1 คันจะมีงานในวันทำงาน
CANCEL_RATE = 0.02             # สัดส่วนงานยกเลิก
AVG_SPEED_KMH = 50             # ความเร็วเฉลี่ยรวมรถติด
LOAD_HOURS = 1.0               # เวลาขึ้นของที่คลัง
UNLOAD_HOURS = 0.5             # เวลาลงของที่ปลายทาง
CUSTOMERS = [f"CUST-{i:02d}" for i in range(1, 9)]
CUSTOMER_DISCOUNT = {c: d for c, d in zip(CUSTOMERS, [0, 0, 0.03, 0.05, 0, 0.02, 0.08, 0])}

# =====================================================================
# 1) Vehicle master
# =====================================================================
vehicles, idx = [], 1
for vt, p in VEHICLE_TYPES.items():
    for k in range(p["n"]):
        status = "Standby" if k >= p["n"] - p["standby"] else "Active"
        vehicles.append(dict(
            Vehicle_Code=f"VEH-{idx:02d}", Vehicle_Type=vt, Fuel_Type=FUEL_TYPE,
            Status=status, Purchase_Date=date(2022 + (idx % 3), 1 + (idx % 12), 1),
            Purchase_Price_THB=p["price"], Useful_Life_Years=p["life"],
            Salvage_Value_THB=round(p["price"] * p["salvage"]),
            Max_Payload_kg=p["capacity"], Rated_km_per_L=p["kmpl"],
        ))
        idx += 1
veh_df = pd.DataFrame(vehicles)

# =====================================================================
# 2) Route master  (ไป-กลับ เพราะรถต้องตีเปล่ากลับคลัง)
# =====================================================================
route_df = pd.DataFrame([dict(
    Route_Code=f"{DEPOT}-{c}", Origin=DEPOT, Destination=c,
    Destination_Province_TH=th, Destination_Province_EN=en, Delivery_Area=area,
    One_Way_km=km, Round_Trip_km=km * 2,
    Toll_RoundTrip_4W_THB=t4, Toll_RoundTrip_6W_THB=t6, Toll_RoundTrip_10W_THB=t10,
) for c, th, en, area, km, t4, t6, t10 in ROUTES])

# =====================================================================
# 3) Trips  — ข้อมูลดิบรายเที่ยว
# =====================================================================
def workdays(a, b):
    d = a
    while d <= b:
        if d.weekday() < 6:          # จันทร์–เสาร์
            yield d
        d += timedelta(days=1)

def fmt(dt):
    return dt.strftime("%Y-%m-%d %H:%M")

trips, n = [], 1
active = veh_df[veh_df.Status == "Active"]
for d in workdays(START, END):
    for _, v in active.iterrows():
        if random.random() > TRIP_PROB_PER_WORKDAY:
            continue
        p = VEHICLE_TYPES[v.Vehicle_Type]
        r = route_df.sample(1, random_state=random.randint(0, 10**9)).iloc[0]
        cust = random.choice(CUSTOMERS)
        cancelled = random.random() < CANCEL_RATE

        dep = datetime.combine(d, datetime.min.time()) + timedelta(
            hours=random.randint(6, 13), minutes=random.choice([0, 15, 30, 45]))
        plan_h = LOAD_HOURS + r.One_Way_km / AVG_SPEED_KMH
        sched = dep + timedelta(hours=plan_h + 0.5)                     # เผื่อเวลา 30 นาที
        actual = dep + timedelta(hours=plan_h * random.lognormvariate(0, 0.15))
        done = actual + timedelta(hours=UNLOAD_HOURS + r.One_Way_km / AVG_SPEED_KMH)

        # ราคา = (ค่าบริการเริ่มต้น + อัตรา/กม. × ระยะเที่ยวเดียว) × (1 − ส่วนลดลูกค้า)
        price = (p["base_fare"] + p["rate_km"] * r.One_Way_km) * (1 - CUSTOMER_DISCOUNT[cust])
        load = round(p["capacity"] * random.uniform(0.5, 0.95), -1)
        # น้ำมันจริง = ระยะ / km/L × ผลของน้ำหนักบรรทุก × ความแปรปรวนการขับ
        load_factor = 1 + 0.10 * (load / p["capacity"])   # บรรทุกเต็ม → กินน้ำมันเพิ่ม ~10%
        liters = r.Round_Trip_km / p["kmpl"] * load_factor * random.gauss(1, 0.04)

        trips.append(dict(
            Trip_ID=f"TR-{n:05d}", Trip_Date=d, Customer_Code=cust,
            Vehicle_Code=v.Vehicle_Code, Vehicle_Type=v.Vehicle_Type, Route_Code=r.Route_Code,
            Distance_km=0 if cancelled else r.Round_Trip_km,
            Load_kg=0 if cancelled else load,
            Departure_Time=None if cancelled else fmt(dep),
            Scheduled_Delivery=fmt(sched),
            Actual_Delivery=None if cancelled else fmt(actual),
            Completion_Time=None if cancelled else fmt(done),
            Fuel_Liters=0 if cancelled else round(liters, 2),
            Fuel_Price_THB_per_L=DIESEL_PRICE[d.month],
            Toll_THB=0 if cancelled else int(r[f"Toll_RoundTrip_{v.Vehicle_Type}_THB"]),
            Driver_Allowance_THB=0 if cancelled else p["allowance"],
            Agreed_Price_THB=0 if cancelled else round(price, 2),
            Job_Status="Cancelled" if cancelled else "Completed",
        ))
        n += 1
trip_df = pd.DataFrame(trips)

# =====================================================================
# 4) Cost parameters (ต่อประเภทรถ) + Emission factors
# =====================================================================
cost_df = pd.DataFrame([dict(
    Vehicle_Type=vt, Driver_Salary_THB_per_Month=p["salary"],
    Insurance_Tax_THB_per_Year=p["ins_tax"], Tire_Count=p["tires"],
    Tire_Price_THB=p["tire_price"], Tire_Life_km=p["tire_life"],
    Oil_Change_Cost_THB=p["oil_cost"], Oil_Change_Interval_km=p["oil_interval"],
    Repair_THB_per_km=p["repair_km"], Base_Fare_THB=p["base_fare"],
    Rate_THB_per_km_OneWay=p["rate_km"],
) for vt, p in VEHICLE_TYPES.items()])

ef_df = pd.DataFrame([dict(Fuel_Type="Diesel", kgCO2_per_L=EF_DIESEL,
    Source="IPCC 2006 default (74,100 kgCO2/TJ × ~36.4 MJ/L) — ให้แทนด้วยค่าของ อบก. (TGO) ฉบับล่าสุด")])

if __name__ == "__main__":
    print(f"vehicles={len(veh_df)} routes={len(route_df)} trips={len(trip_df)}")

# =====================================================================
# 5) คำอธิบายคอลัมน์ 2 ภาษา + สมมติฐาน → เขียนไฟล์ Excel
# =====================================================================
COLUMN_GUIDE = [
    # sheet, column, ไทย, หน่วย, ที่มา
    ("Trips", "Trip_ID", "รหัสเที่ยวขนส่ง", "-", "รหัสลำดับ"),
    ("Trips", "Trip_Date", "วันที่ขนส่ง", "วันที่", "วันทำงาน จ.–ส."),
    ("Trips", "Customer_Code", "รหัสลูกค้า", "-", "สุ่มจากลูกค้า 8 ราย"),
    ("Trips", "Vehicle_Code", "รหัสรถ", "-", "อ้างอิงชีต Vehicles"),
    ("Trips", "Vehicle_Type", "ประเภทรถ", "4W/6W/10W", "อ้างอิงชีต Vehicles"),
    ("Trips", "Route_Code", "รหัสเส้นทาง", "-", "อ้างอิงชีต Routes"),
    ("Trips", "Distance_km", "ระยะทางไป-กลับ", "กม.", "= Routes.Round_Trip_km (เส้นทางเดียวกันระยะเท่ากันเสมอ)"),
    ("Trips", "Load_kg", "น้ำหนักสินค้า", "กก.", "50–95% ของน้ำหนักบรรทุกสูงสุด"),
    ("Trips", "Departure_Time", "เวลาออกจากคลัง", "วันเวลา", "06:00–13:45"),
    ("Trips", "Scheduled_Delivery", "เวลาส่งตามกำหนด", "วันเวลา", "ออกเดินทาง + เวลาขึ้นของ + ระยะ/ความเร็วเฉลี่ย + 30 นาที"),
    ("Trips", "Actual_Delivery", "เวลาส่งถึงจริง", "วันเวลา", "เวลาวางแผน × ความแปรปรวนการจราจร"),
    ("Trips", "Completion_Time", "เวลากลับถึงคลัง", "วันเวลา", "ส่งถึง + ลงของ + เวลาตีรถเปล่ากลับ"),
    ("Trips", "Fuel_Liters", "ปริมาณน้ำมันที่ใช้", "ลิตร", "ระยะ ÷ km/L × ผลของน้ำหนักบรรทุก × ความแปรปรวน"),
    ("Trips", "Fuel_Price_THB_per_L", "ราคาน้ำมัน", "บาท/ลิตร", "ราคาดีเซลเฉลี่ยรายเดือน (สมมติ)"),
    ("Trips", "Toll_THB", "ค่าทางด่วนไป-กลับ", "บาท", "อ้างอิงชีต Routes ตามประเภทรถ"),
    ("Trips", "Driver_Allowance_THB", "เบี้ยเลี้ยงคนขับ", "บาท/เที่ยว", "อ้างอิงประเภทรถ"),
    ("Trips", "Agreed_Price_THB", "ค่าขนส่งที่ตกลงกับลูกค้า (รายได้)", "บาท", "(ค่าเริ่มต้น + อัตรา/กม. × ระยะเที่ยวเดียว) × (1 − ส่วนลดลูกค้า)"),
    ("Trips", "Job_Status", "สถานะงาน", "-", "Cancelled ไม่มีรายได้/ระยะทาง/น้ำมัน"),
    ("Vehicles", "Status", "สถานะรถ", "Active/Standby", "Standby = รถจอดสำรอง ยังต้องคิดค่าเสื่อม"),
    ("Vehicles", "Purchase_Price_THB", "ราคาซื้อรถ", "บาท", "ใช้คิดค่าเสื่อมราคา"),
    ("Vehicles", "Useful_Life_Years", "อายุการใช้งาน", "ปี", "ใช้คิดค่าเสื่อมราคาแบบเส้นตรง"),
    ("Vehicles", "Salvage_Value_THB", "มูลค่าซาก", "บาท", "10% ของราคาซื้อ (สมมติ)"),
    ("Vehicles", "Max_Payload_kg", "น้ำหนักบรรทุกสูงสุด", "กก.", ""),
    ("Vehicles", "Rated_km_per_L", "อัตราสิ้นเปลืองมาตรฐาน", "กม./ลิตร", "ใช้สร้างข้อมูลน้ำมัน"),
    ("Routes", "Destination_Province_TH", "จังหวัดปลายทาง", "-", "ความหมายของรหัสปลายทาง"),
    ("Routes", "One_Way_km", "ระยะทางเที่ยวเดียว", "กม.", "ระยะถนนโดยประมาณ ต้องตรวจกับ Google Maps"),
    ("Routes", "Round_Trip_km", "ระยะทางไป-กลับ", "กม.", "= One_Way_km × 2"),
]

ASSUMPTIONS = [
    ("วิธีสร้างข้อมูล", "สร้างด้วยสคริปต์ Python (generate_mockup.py) จากสมมติฐานในชีตนี้ ใช้ random seed = 2026 รันซ้ำได้ผลเดิม ไม่ได้ให้ AI แต่งตัวเลขโดยตรง"),
    ("ช่วงข้อมูล", "1 เม.ย. – 30 ก.ย. 2026 (6 เดือน) วันทำงาน จันทร์–เสาร์"),
    ("กองรถ", "20 คัน: 4W 7 คัน (สำรอง 1), 6W 7 คัน (สำรอง 1), 10W 6 คัน; รถสำรองไม่มีเที่ยววิ่งแต่ยังมีค่าเสื่อม/ประกัน"),
    ("ความถี่งาน", "รถ Active แต่ละคันมีโอกาสได้งาน 85% ต่อวันทำงาน, วันละไม่เกิน 1 เที่ยว"),
    ("ระยะทาง", "รถออกจากคลังกรุงเทพฯ ไปส่งแล้วตีเปล่ากลับ → ระยะต่อเที่ยว = ไป-กลับ; เส้นทางเดียวกันระยะเท่ากันเสมอ"),
    ("อัตราสิ้นเปลือง", "4W 10, 6W 6, 10W 4 กม./ลิตร (ค่าทั่วไปโดยประมาณ ต้องหาเอกสารอ้างอิง) บรรทุกเต็มกินน้ำมันเพิ่ม ~10%"),
    ("ราคาค่าขนส่ง", "อัตราแบบ ค่าเริ่มต้น + บาท/กม.: 4W 1,000+24, 6W 1,700+35, 10W 2,700+48 (คิดจากระยะเที่ยวเดียว) ลูกค้าบางรายได้ส่วนลด 2–8%"),
    ("งานยกเลิก", "~2% ของงาน ไม่มีรายได้ ระยะทาง และค่าน้ำมัน"),
    ("ราคาน้ำมัน", "ดีเซลเฉลี่ยรายเดือน 31.9–32.9 บาท/ลิตร (สมมติ) ให้แทนด้วยราคาจริงจาก สนพ."),
    ("ค่าเสื่อมราคา", "ไม่อยู่ในไฟล์เที่ยววิ่ง เว็บคำนวณเองแบบเส้นตรงจากชีต Vehicles: (ราคาซื้อ − ซาก) ÷ อายุใช้งาน"),
    ("ต้นทุนอื่น", "เงินเดือนคนขับ ประกัน/ภาษี ค่ายาง ค่าน้ำมันเครื่อง ค่าซ่อม อยู่ในชีต Cost_Parameters เว็บเป็นผู้คำนวณ"),
    ("คาร์บอน", "ลิตรน้ำมัน × Emission Factor ตามชนิดเชื้อเพลิง (ชีต Emission_Factors)"),
    ("ข้อจำกัด", "เป็นข้อมูลจำลอง ไม่ใช่ข้อมูลบริษัท ตัวเลขทุกตัวต้องแทนด้วยค่าจริงหรือค่าที่มีแหล่งอ้างอิงก่อนนำไปสรุปผล"),
]


def write_excel(path="Transport_Mockup_v2.xlsx"):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    guide_df = pd.DataFrame(COLUMN_GUIDE, columns=["Sheet", "Column (EN)", "ชื่อภาษาไทย", "หน่วย", "ที่มา / สูตรที่ใช้สร้าง"])
    assump_df = pd.DataFrame(ASSUMPTIONS, columns=["หัวข้อ", "สมมติฐาน"])
    sheets = {
        "Read_Me": assump_df, "Trips": trip_df, "Vehicles": veh_df, "Routes": route_df,
        "Cost_Parameters": cost_df, "Emission_Factors": ef_df, "Column_Guide": guide_df,
    }
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)
        head = PatternFill("solid", start_color="1F4E78")
        blue = Font(name="Arial", color="0000FF")
        for name, ws in zip(xw.book.sheetnames, xw.book.worksheets):
            for row in ws.iter_rows():
                for c in row:
                    c.font = Font(name="Arial")
            for c in ws[1]:
                c.font = Font(name="Arial", bold=True, color="FFFFFF")
                c.fill = head
                c.alignment = Alignment(wrap_text=True, vertical="center")
            # ชีตพารามิเตอร์ = ค่าที่ป้อนเอง → ตัวอักษรสีน้ำเงิน (ตามหลัก input cell)
            if name in ("Vehicles", "Routes", "Cost_Parameters", "Emission_Factors"):
                for row in ws.iter_rows(min_row=2):
                    for c in row:
                        if isinstance(c.value, (int, float)):
                            c.font = blue
            for i, col in enumerate(sheets[name].columns, 1):
                width = max(len(str(col)), *(len(str(v)) for v in sheets[name][col].head(50))) + 2
                ws.column_dimensions[get_column_letter(i)].width = min(width, 80)
            ws.freeze_panes = "A2"
    print("saved", path)


if __name__ == "__main__":
    write_excel()
